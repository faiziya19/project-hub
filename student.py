import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


FILE_NAME = "student_results.xlsx"

columns = [
    "Name", "Roll No.", "Class",
    "Subject 1", "Subject 2", "Subject 3",
    "Subject 4", "Subject 5",
    "Total Marks", "Percentage", "Result"
]


if not os.path.exists(FILE_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(columns)
    wb.save(FILE_NAME)


root = tk.Tk()
root.title("Student Result Management System")
root.geometry("900x600")
root.resizable(False, False)


main_frame = tk.Frame(root)
main_frame.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=20
)



def clear_screen():

    for widget in main_frame.winfo_children():
        widget.destroy()






def create_title(title):

    tk.Label(
        main_frame,
        text=title,
        font=("Arial", 24, "bold")
    ).pack(pady=20)






def add_student():

    clear_screen()

    create_title("1. Add Student")

    form = tk.Frame(main_frame)
    form.pack(pady=10)

    labels = [
        "Student Name",
        "Roll No.",
        "Class",
        "Subject 1 Marks",
        "Subject 2 Marks",
        "Subject 3 Marks",
        "Subject 4 Marks",
        "Subject 5 Marks"
    ]

    entries = []


    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label + ":",
            font=("Arial", 11)
        ).grid(
            row=i,
            column=0,
            padx=15,
            pady=6,
            sticky="w"
        )

        entry = tk.Entry(
            form,
            width=30,
            font=("Arial", 11)
        )

        entry.grid(
            row=i,
            column=1,
            padx=15,
            pady=6
        )

        entries.append(entry)


    def save_student():

        name = entries[0].get().strip()
        roll = entries[1].get().strip()
        student_class = entries[2].get().strip()



        if name == "" or roll == "" or student_class == "":
            messagebox.showerror(
                "Error",
                "Please enter Name, Roll No. and Class."
            )
            return


        try:
            roll_number = int(roll)

        except ValueError:
            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )
            return


        marks = []

        try:

            for i in range(3, 8):

                mark = float(entries[i].get())

                if mark < 0 or mark > 100:

                    messagebox.showerror(
                        "Error",
                        "Marks must be between 0 and 100."
                    )

                    return

                marks.append(mark)

        except ValueError:

            messagebox.showerror(
                "Error",
                "Please enter valid marks."
            )

            return


        wb = load_workbook(FILE_NAME)
        ws = wb.active


        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if row[1] == roll_number:

                wb.close()

                messagebox.showerror(
                    "Error",
                    "This Roll No. already exists."
                )

                return


        total = sum(marks)


        percentage = total / 5

 
        if all(mark >= 40 for mark in marks):
            result = "Pass"

        else:
            result = "Fail"
 

        ws.append([
            name,
            roll_number,
            student_class,
            *marks,
            total,
            percentage,
            result
        ])

        wb.save(FILE_NAME)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Student record saved successfully!"
        )



     

        get_result()





    tk.Button(
        main_frame,
        text="💾 SAVE",
        font=("Arial", 13, "bold"),
        width=18,
        command=save_student
    ).pack(pady=20)






def get_result():

    clear_screen()

    create_title("2. Get Result")

    tk.Label(
        main_frame,
        text="Enter Roll No.:",
        font=("Arial", 13)
    ).pack(pady=10)

    roll_entry = tk.Entry(
        main_frame,
        width=25,
        font=("Arial", 13)
    )

    roll_entry.pack(pady=5)

    result_frame = tk.Frame(main_frame)
    result_frame.pack(pady=20)




    def search_result():

        roll = roll_entry.get().strip()

        if roll == "":

            messagebox.showerror(
                "Error",
                "Please enter Roll No."
            )

            return

        try:

            roll_number = int(roll)

        except ValueError:

            messagebox.showerror(
                "Error",
                "Roll No. must be a number."
            )

            return



        for widget in result_frame.winfo_children():
            widget.destroy()

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

 
        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if row[1] == roll_number:

                found = True

              
                result_data = [
                    ("Name", row[0]),
                    ("Roll No.", row[1]),
                    ("Class", row[2]),
                    ("Total Marks", row[8]),
                    ("Percentage", f"{row[9]:.1f}%"),
                    ("Result", row[10])
                ]

        
                for i, (label, value) in enumerate(result_data):

                    tk.Label(
                        result_frame,
                        text=label + ":",
                        font=("Arial", 11, "bold"),
                        width=15,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=0,
                        padx=10,
                        pady=5
                    )

                    tk.Label(
                        result_frame,
                        text=value,
                        font=("Arial", 11),
                        width=20,
                        anchor="w"
                    ).grid(
                        row=i,
                        column=1,
                        padx=10,
                        pady=5
                    )

                break

        wb.close()

       

        if not found:

            tk.Label(
                result_frame,
                text="❌ Student record not found.",
                font=("Arial", 13, "bold")
            ).pack(pady=20)

            return

       

        tk.Button(
            main_frame,
            text="SUBMIT & CONTINUE",
            font=("Arial", 12, "bold"),
            width=22,
            command=show_all_results
        ).pack(pady=20)



    

    tk.Button(
        main_frame,
        text="🔍 GET RESULT",
        font=("Arial", 12, "bold"),
        width=18,
        command=search_result
    ).pack(pady=10)





def show_all_results():

    clear_screen()

    create_title("3. Show All Results")
 

    table_frame = tk.Frame(main_frame)

    table_frame.pack(
        fill="both",
        expand=True,
        pady=10
    )

    display_columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )


    tree = ttk.Treeview(
        table_frame,
        columns=display_columns,
        show="headings",
        height=14
    )


    
    for column in display_columns:

        tree.heading(
            column,
            text=column
        )

        tree.column(
            column,
            width=130,
            anchor="center"
        )

    tree.pack(
        side="left",
        fill="both",
        expand=True
    )
    

    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    tree.configure(
        yscrollcommand=scrollbar.set
    )

    scrollbar.pack(
        side="right",
        fill="y"
    )


    wb = load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is not None:

            tree.insert(
                "",
                "end",
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[8],
                    f"{row[9]:.1f}%",
                    row[10]
                )
            )

    wb.close()
 
    button_frame = tk.Frame(main_frame)
    button_frame.pack(pady=15)

    tk.Button(
        button_frame,
        text="➕ ADD ANOTHER STUDENT",
        font=("Arial", 12, "bold"),
        width=25,
        command=add_student
    ).grid(
        row=0,
        column=0,
        padx=10
    )

    tk.Button(
        button_frame,
        text="EXIT",
        font=("Arial", 12, "bold"),
        width=15,
        command=root.destroy
    ).grid(
        row=0,
        column=1,
        padx=10
    )

add_student()

root.mainloop()
